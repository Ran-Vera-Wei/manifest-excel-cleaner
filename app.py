from __future__ import annotations

import io
import re
import zipfile
from bisect import bisect_left
from copy import copy
from pathlib import Path

import msoffcrypto
import streamlit as st
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.utils import get_column_letter, range_boundaries


DESTINATION_CODES = ("EZE", "GUA", "SAL", "LIM", "SJO", "VCP", "SCL", "POS", "SAP", "UIO")
OLE_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")
SENSITIVE_TERMS = (
    "ALCOHOL",
    "ACETONE",
    "ESSENTIAL OIL",
    "PERFUME",
    "PERFUMING",
    "NAIL ART",
    "NAIL POLISH",
)
TERM_PATTERNS = tuple(
    (
        term,
        re.compile(
            r"(?<![A-Z0-9])" + r"[\s\u00a0]+".join(map(re.escape, term.split())) + r"(?![A-Z0-9])",
            re.IGNORECASE,
        ),
    )
    for term in SENSITIVE_TERMS
)


class ProcessingError(Exception):
    """An expected, user-correctable workbook or input error."""


def password_mapping() -> dict[str, str]:
    """Read destination passwords from Streamlit Secrets, never from source code."""
    if "manifest_passwords" not in st.secrets:
        return {}
    try:
        section = st.secrets["manifest_passwords"]
        return {str(code).strip().upper(): str(password) for code, password in section.items()}
    except (KeyError, AttributeError, TypeError):
        return {}


def decrypt_if_needed(file_bytes: bytes, destination_code: str, passwords: dict[str, str]) -> bytes:
    if file_bytes.startswith(OLE_SIGNATURE):
        if destination_code not in DESTINATION_CODES:
            raise ProcessingError("This file requires an opening password. Enter the matching three-character destination code, such as SAL.")
        password = passwords.get(destination_code)
        if not password:
            raise ProcessingError(
                "No password is configured for this destination. Add its password to the [manifest_passwords] section in Streamlit Secrets and try again."
            )

        try:
            encrypted = msoffcrypto.OfficeFile(io.BytesIO(file_bytes))
            if not encrypted.is_encrypted():
                raise ProcessingError("This file is not a supported encrypted Excel workbook.")
            encrypted.load_key(password=password, verify_password=True)
            decrypted = io.BytesIO()
            encrypted.decrypt(decrypted)
            return decrypted.getvalue()
        except ProcessingError:
            raise
        except Exception as exc:
            raise ProcessingError("Unable to open the file. Check the three-character destination code and confirm its password is configured in Secrets.") from exc

    if not zipfile.is_zipfile(io.BytesIO(file_bytes)):
        raise ProcessingError("Only .xlsx and .xlsm files are supported. Upload a valid Excel workbook.")
    return file_bytes


def normalized_header(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip().casefold()


def find_manifest_header(ws) -> int | None:
    for row_number in range(1, min(ws.max_row, 100) + 1):
        bg_header = normalized_header(ws.cell(row=row_number, column=3).value)
        description_header = normalized_header(ws.cell(row=row_number, column=15).value)
        if bg_header == "bg number" and description_header == "product description":
            return row_number
    return None


def remove_duplicate_bg_rows(ws, header_row: int) -> int:
    seen: set[str] = set()
    rows_to_remove: list[int] = []
    for row_number in range(header_row + 1, ws.max_row + 1):
        value = ws.cell(row=row_number, column=3).value
        if value is None or not str(value).strip():
            continue
        key = str(value).strip().casefold()
        if key in seen:
            rows_to_remove.append(row_number)
        else:
            seen.add(key)

    if not rows_to_remove:
        return 0

    if any(cell.data_type == "f" and cell.row > header_row for cell in ws._cells.values()):
        raise ProcessingError(
            f"Worksheet '{ws.title}' contains formulas in its data area. Convert them to values before processing to avoid shifting formula references."
        )
    if any(merged.max_row > header_row for merged in ws.merged_cells.ranges):
        raise ProcessingError(
            f"Worksheet '{ws.title}' contains merged cells in its data area, so duplicate rows cannot be removed safely."
        )
    if ws.conditional_formatting or ws.data_validations.dataValidation:
        raise ProcessingError(
            f"Worksheet '{ws.title}' contains conditional formatting or data validation. Remove these rules before processing to prevent their references from shifting."
        )

    def remap_row(row_number: int) -> int | None:
        removed_before = bisect_left(rows_to_remove, row_number)
        if removed_before < len(rows_to_remove) and rows_to_remove[removed_before] == row_number:
            return None
        return row_number - removed_before if row_number > header_row else row_number

    # Compact the worksheet once instead of shifting every later row once per duplicate.
    compacted_cells = {}
    for (old_row, column), cell in ws._cells.items():
        new_row = remap_row(old_row)
        if new_row is None:
            continue
        if new_row != old_row:
            cell.row = new_row
        compacted_cells[(new_row, column)] = cell
    ws._cells = compacted_cells

    old_dimensions = list(ws.row_dimensions.items())
    ws.row_dimensions.clear()
    for old_row, dimension in old_dimensions:
        new_row = remap_row(old_row)
        if new_row is None:
            continue
        dimension.index = new_row
        ws.row_dimensions[new_row] = dimension

    if ws.auto_filter.ref:
        min_col, min_row, max_col, max_row = range_boundaries(ws.auto_filter.ref)
        new_min_row = remap_row(min_row)
        new_max_row = remap_row(max_row)
        if new_min_row is not None and new_max_row is not None:
            ws.auto_filter.ref = (
                f"{get_column_letter(min_col)}{new_min_row}:"
                f"{get_column_letter(max_col)}{new_max_row}"
            )
    for table in ws.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        new_min_row = remap_row(min_row)
        new_max_row = remap_row(max_row)
        if new_min_row is not None and new_max_row is not None:
            table.ref = (
                f"{get_column_letter(min_col)}{new_min_row}:"
                f"{get_column_letter(max_col)}{new_max_row}"
            )
            if table.autoFilter is not None:
                table.autoFilter.ref = table.ref
    return len(rows_to_remove)


def remove_sensitive_terms(ws, header_row: int) -> dict[str, int]:
    counts = {term: 0 for term in SENSITIVE_TERMS}
    for row_number in range(header_row + 1, ws.max_row + 1):
        cell = ws.cell(row=row_number, column=15)
        if not isinstance(cell.value, str) or cell.data_type == "f":
            continue
        cleaned = cell.value
        for term, pattern in TERM_PATTERNS:
            matches = pattern.findall(cleaned)
            counts[term] += len(matches)
            cleaned = pattern.sub("", cleaned)
        cleaned = re.sub(r"[ \t]{2,}", " ", cleaned).strip()
        cell.value = cleaned
    return counts


def apply_font_size_and_print_settings(workbook) -> None:
    font_ids = {}
    for old_font_id, font in enumerate(tuple(workbook._fonts)):
        updated_font = copy(font)
        updated_font.sz = 16
        font_ids[old_font_id] = workbook._fonts.add(updated_font)

    for ws in workbook.worksheets:
        updated_styles = {}
        default_style = workbook._cell_styles[0]
        for cell in ws._cells.values():
            if isinstance(cell, MergedCell):
                continue
            original_style = cell._style or default_style
            style_key = tuple(original_style)
            updated_style = updated_styles.get(style_key)
            if updated_style is None:
                updated_style = copy(original_style)
                updated_style.fontId = font_ids[original_style.fontId]
                updated_styles[style_key] = updated_style
            cell._style = updated_style

        setup = ws.sheet_properties.pageSetUpPr
        if setup is None:
            ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True, autoPageBreaks=False)
        else:
            setup.fitToPage = True
            setup.autoPageBreaks = False

        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.75
        ws.page_margins.bottom = 0.75
        ws.page_margins.header = 0.3
        ws.page_margins.footer = 0.3

        if ws.max_row and ws.max_column:
            ws.print_area = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


def process_workbook(
    file_bytes: bytes,
    file_name: str,
    destination_code: str,
    passwords: dict[str, str],
) -> tuple[bytes, str, dict[str, dict[str, object]]]:
    suffix = Path(file_name).suffix.lower()
    if suffix not in (".xlsx", ".xlsm"):
        raise ProcessingError("Only .xlsx and .xlsm files are supported.")

    workbook_bytes = decrypt_if_needed(file_bytes, destination_code, passwords)
    try:
        workbook = load_workbook(
            io.BytesIO(workbook_bytes),
            data_only=False,
            keep_vba=suffix == ".xlsm",
        )
    except Exception as exc:
        raise ProcessingError("Unable to read the Excel workbook. Confirm that the file is not damaged and is in .xlsx or .xlsm format.") from exc

    summaries: dict[str, dict[str, object]] = {}
    matching_sheets = []
    for ws in workbook.worksheets:
        header_row = find_manifest_header(ws)
        if header_row is None:
            continue
        matching_sheets.append((ws, header_row))

    if not matching_sheets:
        raise ProcessingError(
            "Could not find the expected headers. Confirm that a row has 'BG Number' in column C and 'PRODUCT DESCRIPTION' in column O."
        )

    for ws, header_row in matching_sheets:
        duplicates_removed = remove_duplicate_bg_rows(ws, header_row)
        terms_removed = remove_sensitive_terms(ws, header_row)
        summaries[ws.title] = {
            "Duplicate rows removed": duplicates_removed,
            "Sensitive term removals": terms_removed,
        }

    apply_font_size_and_print_settings(workbook)
    try:
        if hasattr(workbook, "calculation"):
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
            workbook.calculation.calcMode = "auto"
    except (AttributeError, TypeError):
        pass

    output = io.BytesIO()
    try:
        workbook.save(output)
    except Exception as exc:
        raise ProcessingError("Unable to save the processed Excel file. Confirm that the file is not damaged and try again.") from exc

    output_bytes = output.getvalue()
    extension = ".xlsm" if suffix == ".xlsm" else ".xlsx"
    return output_bytes, extension, summaries


def safe_output_name(awb: str, extension: str) -> str:
    safe_awb = re.sub(r"[^A-Za-z0-9._-]+", "", awb.strip()).strip("._-")
    if not safe_awb:
        raise ProcessingError("Enter a valid export AWB#.")
    return f"{safe_awb}manifest{extension}"


def main() -> None:
    st.set_page_config(page_title="Manifest Excel Cleaner", page_icon="📄", layout="centered")
    st.title("Manifest Excel Cleaner")
    st.write("Upload an Excel file, unlock it with a destination code if needed, remove duplicate BG Numbers and sensitive terms from product descriptions, then download the result.")

    uploaded_file = st.file_uploader("Upload a manifest (.xlsx or .xlsm)", type=["xlsx", "xlsm"])
    destination_code = st.text_input("Three-character destination code", max_chars=3, placeholder="e.g. SAL").strip().upper()
    awb = st.text_input("Export AWB#", placeholder="e.g. 810-51887706").strip()

    if st.button("Process and generate file", type="primary", use_container_width=True):
        if uploaded_file is None:
            st.error("Upload a manifest file first.")
        elif not awb:
            st.error("Enter an export AWB#.")
        elif destination_code and len(destination_code) != 3:
            st.error("The destination code must contain exactly three characters.")
        else:
            try:
                output_bytes, extension, summaries = process_workbook(
                    uploaded_file.getvalue(),
                    uploaded_file.name,
                    destination_code,
                    password_mapping(),
                )
                output_name = safe_output_name(awb, extension)
                st.success("Processing complete. The output workbook is not protected with an opening password.")
                for sheet_name, summary in summaries.items():
                    st.subheader(sheet_name)
                    st.write(f"Duplicate rows removed: {summary['Duplicate rows removed']}")
                    term_counts = summary["Sensitive term removals"]
                    st.write("Sensitive term removals: " + "; ".join(f"{term}: {count}" for term, count in term_counts.items()))
                st.download_button(
                    label="Download processed manifest",
                    data=output_bytes,
                    file_name=output_name,
                    mime=(
                        "application/vnd.ms-excel.sheet.macroEnabled.12"
                        if extension == ".xlsm"
                        else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    ),
                    use_container_width=True,
                )
            except ProcessingError as exc:
                st.error(str(exc))
            except Exception:
                st.error("Processing failed. Check the file and your inputs, then try again.")

    st.caption(
        "Print settings: narrow margins, landscape orientation, all columns fit on one page, and 16 pt font. "
        "Duplex printing is controlled by the printer driver; select two-sided printing in the print dialog."
    )


if __name__ == "__main__":
    main()
